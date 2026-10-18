import os
import re
from typing import Any, Dict

# Directory where notes and documents are saved safely
NOTES_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "notes"))


class ToolValidationError(Exception):
    """Raised when tool inputs or outputs fail validation."""
    pass


def _sanitize_title(title: Any) -> str:
    """Validates and sanitizes a file or document title."""
    if not isinstance(title, str):
        raise ToolValidationError(f"Invalid 'title': expected string, got {type(title).__name__}")
    clean_title = title.strip()
    if not clean_title:
        raise ToolValidationError("Title cannot be empty.")
    if len(clean_title) > 100:
        raise ToolValidationError(f"Title is too long ({len(clean_title)} chars, max 100).")
    # Replace dangerous characters with underscore
    clean_title = re.sub(r'[\/\\:\*\?"<>\|\0\.]+', '_', clean_title).strip('_')
    if not clean_title:
        clean_title = "Document"
    return clean_title


def validate_save_note_inputs(title: Any, content: Any, category: Any = "general") -> Dict[str, str]:
    """Validates inputs for the `save_note` tool."""
    clean_title = _sanitize_title(title)

    if not isinstance(content, str):
        raise ToolValidationError(f"Invalid 'content': expected string, got {type(content).__name__}")
    clean_content = content.strip()
    if not clean_content:
        raise ToolValidationError("Note content cannot be empty.")
    if len(clean_content) > 100000:
        raise ToolValidationError(f"Content is too long ({len(clean_content)} chars, max 100,000).")

    clean_category = (category or "general").strip().lower()
    clean_category = re.sub(r'[^a-zA-Z0-9_-]', '_', clean_category)[:50] or "general"

    return {
        "title": clean_title,
        "content": clean_content,
        "category": clean_category,
    }


def validate_save_note_output(output: Dict[str, Any]) -> None:
    """Validates that the tool output conforms to expected contract."""
    if not isinstance(output, dict):
        raise ToolValidationError(f"Tool output must be a dict, got {type(output).__name__}")
    if output.get("status") != "success":
        raise ToolValidationError(f"Tool did not report success: {output.get('message', 'Unknown error')}")
    file_path = output.get("file_path")
    if not file_path or not os.path.exists(file_path):
        raise ToolValidationError(f"Output validation failed: file '{file_path}' was not found on disk.")
    if os.path.getsize(file_path) == 0:
        raise ToolValidationError(f"Output validation failed: file '{file_path}' is empty.")


def save_note(title: str, content: str, category: str = "general", base_dir: str = None) -> Dict[str, Any]:
    """Saves a note to the filesystem as a text or markdown file."""
    try:
        validated = validate_save_note_inputs(title=title, content=content, category=category)
        target_dir = os.path.abspath(base_dir) if base_dir else NOTES_DIR
        os.makedirs(target_dir, exist_ok=True)

        sanitized_filename = f"{validated['category']}_{validated['title']}.md"
        file_path = os.path.abspath(os.path.join(target_dir, sanitized_filename))

        if not file_path.startswith(target_dir):
            raise ToolValidationError("Path traversal detected: target path is outside notes directory.")

        file_content = f"# {validated['title']}\nCategory: {validated['category']}\n\n{validated['content']}\n"
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(file_content)

        bytes_written = os.path.getsize(file_path)

        result = {
            "status": "success",
            "message": f"Note '{validated['title']}' saved successfully in category '{validated['category']}'.",
            "file_path": file_path,
            "filename": sanitized_filename,
            "doc_type": "md",
            "bytes_written": bytes_written,
        }
        validate_save_note_output(result)
        return result
    except ToolValidationError as ve:
        return {"status": "error", "error_type": "ValidationError", "message": str(ve)}
    except Exception as e:
        return {"status": "error", "error_type": "ExecutionError", "message": str(e)}


def create_document(title: str, content: str, doc_type: str = "docx", base_dir: str = None) -> Dict[str, Any]:
    """
    Creates and saves a structured document (Word .docx, Markdown .md, or text)
    with professional formatting, headings, bullet points, and clean margins.
    """
    try:
        clean_title = _sanitize_title(title)
        clean_content = content.strip() if isinstance(content, str) else ""
        if not clean_content:
            raise ToolValidationError("Document content cannot be empty.")

        clean_type = (doc_type or "docx").strip().lower().replace(".", "")
        if clean_type not in ("docx", "md", "txt"):
            clean_type = "docx"

        target_dir = os.path.abspath(base_dir) if base_dir else NOTES_DIR
        os.makedirs(target_dir, exist_ok=True)

        filename = f"{clean_title}.{clean_type}"
        file_path = os.path.abspath(os.path.join(target_dir, filename))

        if not file_path.startswith(target_dir):
            raise ToolValidationError("Path traversal detected.")

        if clean_type == "docx":
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()
            # 1 inch margins
            for s in doc.sections:
                s.top_margin = Inches(1)
                s.bottom_margin = Inches(1)
                s.left_margin = Inches(1)
                s.right_margin = Inches(1)

            # Main Title
            title_p = doc.add_heading(clean_title, level=0)
            title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            lines = clean_content.split("\n")
            for raw_line in lines:
                line = raw_line.strip()
                if not line:
                    continue
                if line.startswith("### "):
                    doc.add_heading(line[4:].strip(), level=3)
                elif line.startswith("## "):
                    doc.add_heading(line[3:].strip(), level=2)
                elif line.startswith("# "):
                    doc.add_heading(line[2:].strip(), level=1)
                elif line.startswith("- ") or line.startswith("* "):
                    doc.add_paragraph(line[2:].strip(), style="List Bullet")
                elif line[0].isdigit() and len(line) > 2 and line[1] in (".", ")"):
                    doc.add_paragraph(line[2:].strip(), style="List Number")
                else:
                    doc.add_paragraph(line)

            doc.save(file_path)
        else:
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(f"# {clean_title}\n\n{clean_content}\n")

        bytes_written = os.path.getsize(file_path)

        result = {
            "status": "success",
            "message": f"Document '{filename}' created successfully.",
            "file_path": file_path,
            "filename": filename,
            "doc_type": clean_type,
            "bytes_written": bytes_written,
        }
        validate_save_note_output(result)
        return result
    except ToolValidationError as ve:
        return {"status": "error", "error_type": "ValidationError", "message": str(ve)}
    except Exception as e:
        return {"status": "error", "error_type": "ExecutionError", "message": str(e)}


# ----------------- Gemini Tool Declarations -----------------
SAVE_NOTE_DECLARATION = {
    "name": "save_note",
    "description": (
        "Saves a quick thought, note, reminder, or snippet to the filesystem. "
        "Use this tool when the user asks to save a note, record a thought, or keep a note."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "title": {
                "type": "STRING",
                "description": "Short title for the note (e.g. 'meeting_ideas').",
            },
            "content": {
                "type": "STRING",
                "description": "The full text content to save in the note.",
            },
            "category": {
                "type": "STRING",
                "description": "Optional category tag (e.g. 'work', 'study', 'ideas', or 'general').",
            },
        },
        "required": ["title", "content"],
    },
}

CREATE_DOCUMENT_DECLARATION = {
    "name": "create_document",
    "description": (
        "Creates, formats, and saves a complete document (such as a Microsoft Word .docx file, "
        "report, essay, article, or guide) on any topic requested by the user, making it ready "
        "for instant download. ALWAYS call this tool when the user asks to make, create, generate, "
        "or write a Word document, docx, report, paper, or file on a topic."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "title": {
                "type": "STRING",
                "description": "Clean, descriptive title of the document (e.g. 'Home' or 'Future_of_AI').",
            },
            "content": {
                "type": "STRING",
                "description": "The complete, rich, well-structured text of the document with markdown headings (#, ##), bullet points, and paragraphs.",
            },
            "doc_type": {
                "type": "STRING",
                "description": "Document format: 'docx' for Microsoft Word document (default), or 'md' for Markdown.",
            },
        },
        "required": ["title", "content"],
    },
}

# ----------------- OpenAI Tool Declarations -----------------
OPENAI_SAVE_NOTE_TOOL = {
    "type": "function",
    "function": {
        "name": "save_note",
        "description": SAVE_NOTE_DECLARATION["description"],
        "parameters": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Short title for the note."},
                "content": {"type": "string", "description": "Full text content."},
                "category": {"type": "string", "description": "Optional category tag."},
            },
            "required": ["title", "content"],
        },
    },
}

OPENAI_CREATE_DOCUMENT_TOOL = {
    "type": "function",
    "function": {
        "name": "create_document",
        "description": CREATE_DOCUMENT_DECLARATION["description"],
        "parameters": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Clean title of the document."},
                "content": {"type": "string", "description": "Complete structured content with headings and bullets."},
                "doc_type": {"type": "string", "description": "Document format: 'docx' or 'md'."},
            },
            "required": ["title", "content"],
        },
    },
}


